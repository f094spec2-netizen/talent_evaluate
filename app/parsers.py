import csv
import io
import json
import zipfile
from pathlib import PurePath

from bs4 import BeautifulSoup
from defusedxml import ElementTree
from openpyxl import load_workbook

SUPPORTED = {".html", ".htm", ".json", ".csv", ".xlsx", ".xlsm", ".docx"}
MAX_RECORDS = 5000
MAX_CHARS = 1_000_000
MIMES = {
    ".html": {"text/html"},
    ".htm": {"text/html"},
    ".json": {"application/json", "text/plain"},
    ".csv": {"text/csv", "text/plain", "application/vnd.ms-excel"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".xlsm": {"application/vnd.ms-excel.sheet.macroenabled.12"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
}


class InvalidFile(ValueError):
    """Safe, stable error code, with no source content or credentials."""


def preflight(filename: str, size: int | None, mime: str | None, limit: int):
    if not filename or len(filename) > 240 or any(c in filename for c in ("/", "\\", "\x00")):
        raise InvalidFile("INVALID_FILENAME")
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise InvalidFile("FORMAT_NOT_SUPPORTED_YET")
    if size is not None and not 0 < size <= limit:
        raise InvalidFile("EMPTY_OR_OVERSIZED_FILE")
    if mime and mime.lower().split(";")[0] not in MIMES[ext] | {"application/octet-stream"}:
        raise InvalidFile("MIME_EXTENSION_MISMATCH")


def safe_office_zip(content: bytes, ext: str):
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise InvalidFile("INVALID_OFFICE_OR_ENCRYPTED")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > 2000 or sum(i.file_size for i in infos) > 40_000_000:
            raise InvalidFile("ARCHIVE_EXPANSION_LIMIT")
        if any(i.flag_bits & 1 for i in infos):
            raise InvalidFile("ENCRYPTED_FILE")
        expected = "word/document.xml" if ext == ".docx" else "xl/workbook.xml"
        names = archive.namelist()
        if (
            expected not in names
            or "[Content_Types].xml" not in names
            or len(names) != len(set(names))
        ):
            raise InvalidFile("OFFICE_CONTENT_MISMATCH")
        if ext == ".xlsx" and any(n.endswith("vbaProject.bin") for n in names):
            raise InvalidFile("MACRO_EXTENSION_MISMATCH")


def parse_file(filename: str, content: bytes) -> tuple[list[dict], list[str]]:
    ext = PurePath(filename).suffix.lower()
    records, warnings = [], []
    characters = 0

    def add(locator, text):
        nonlocal characters
        text = str(text).strip()
        if not text:
            return
        characters += len(text)
        if len(records) >= MAX_RECORDS or characters > MAX_CHARS or len(text) > 50_000:
            raise InvalidFile("EXTRACTION_LIMIT")
        records.append({"locator": locator[:255], "text": text})

    try:
        if ext in {".xlsx", ".xlsm", ".docx"}:
            safe_office_zip(content, ext)
        if ext in {".xlsx", ".xlsm"}:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
            try:
                cells_seen = 0
                for sheet_index, sheet in enumerate(workbook.worksheets, 1):
                    if (
                        sheet.max_row
                        and sheet.max_row > 20_000
                        or sheet.max_column
                        and sheet.max_column > 256
                    ):
                        raise InvalidFile("EXTRACTION_LIMIT")
                    for row_index, row in enumerate(sheet.iter_rows(), 1):
                        cells_seen += len(row)
                        if cells_seen > 200_000:
                            raise InvalidFile("EXTRACTION_LIMIT")
                        values = []
                        for cell in row:
                            if cell.data_type == "f":
                                warnings.append("FORMULA_NOT_EVALUATED")
                            if cell.value is not None:
                                values.append(str(cell.value))
                        add(f"sheet:{sheet_index}/row:{row_index}", " | ".join(values))
                if ext == ".xlsm":
                    warnings.append("MACROS_NOT_EXECUTED")
            finally:
                workbook.close()
        elif ext == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for index, paragraph in enumerate(root.findall(".//w:p", ns), 1):
                add(
                    f"paragraph:{index}",
                    "".join(t.text or "" for t in paragraph.findall(".//w:t", ns)),
                )
        else:
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise InvalidFile("BINARY_CONTENT_MISMATCH")
            if ext in {".html", ".htm"}:
                soup = BeautifulSoup(text, "html.parser")
                if not soup.find(["html", "body", "table", "p", "h1"]):
                    raise InvalidFile("HTML_CONTENT_MISMATCH")
                for tag in soup(["script", "style", "noscript", "iframe", "object"]):
                    tag.decompose()
                for index, line in enumerate(soup.get_text("\n", strip=True).splitlines(), 1):
                    add(f"text_line:{index}", line)
            elif ext == ".csv":
                for index, row in enumerate(csv.reader(io.StringIO(text)), 1):
                    add(f"row:{index}", " | ".join(row))
            elif ext == ".json":
                value = json.loads(text)
                if isinstance(value, dict):
                    for key, item in value.items():
                        add(f"json:/{key}", json.dumps({key: item}, ensure_ascii=False))
                elif isinstance(value, list):
                    for index, item in enumerate(value):
                        add(f"json:/{index}", json.dumps(item, ensure_ascii=False))
                else:
                    raise InvalidFile("JSON_OBJECT_OR_ARRAY_REQUIRED")
            else:
                raise InvalidFile("FORMAT_NOT_SUPPORTED_YET")
    except InvalidFile:
        raise
    except Exception as exc:
        raise InvalidFile("UNREADABLE_OR_DAMAGED_FILE") from exc
    if not records:
        raise InvalidFile("NO_EXTRACTABLE_CONTENT")
    return records, sorted(set(warnings))
